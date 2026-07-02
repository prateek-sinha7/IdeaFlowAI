"""File text extraction endpoint.

Accepts a binary file upload (PDF, DOCX, PPTX) and returns its extracted
plain text, which the frontend appends into the workflow brief so agents
receive the actual document content.

Supported formats:
  - PDF  (.pdf)  — via pypdf
  - Word (.docx) — via python-docx
  - PowerPoint (.pptx) — via python-pptx

Max file size: 10 MB. Extracted text is capped at settings.BRIEF_MAX_CHARS
characters (~450,000) — the single source of truth shared with the planner /
clarify brief cap — to bound the input that reaches the LLM.
"""

from __future__ import annotations

import io
import logging

from fastapi import APIRouter, Depends, HTTPException, UploadFile, status
from pydantic import BaseModel

from app.core.config import settings
from app.core.dependencies import get_current_user
from app.models.user import User

logger = logging.getLogger("app.api.file_extract")

router = APIRouter(prefix="/api/files", tags=["files"])

_MAX_FILE_BYTES = 10 * 1024 * 1024  # 10 MB
# Extracted-upload text cap tracks the brief cap (single source of truth).
_MAX_TEXT_CHARS = settings.BRIEF_MAX_CHARS


class ExtractResponse(BaseModel):
    filename: str
    text: str
    truncated: bool


def _extract_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        parts = []
        for page in reader.pages:
            t = page.extract_text()
            if t:
                parts.append(t)
        return "\n".join(parts)
    except Exception as exc:
        logger.warning("PDF extraction failed: %s", exc)
        return ""


def _extract_docx(data: bytes) -> str:
    try:
        from docx import Document
        doc = Document(io.BytesIO(data))
        return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    except Exception as exc:
        logger.warning("DOCX extraction failed: %s", exc)
        return ""


def _extract_pptx(data: bytes) -> str:
    try:
        from pptx import Presentation
        prs = Presentation(io.BytesIO(data))
        parts = []
        for slide_num, slide in enumerate(prs.slides, 1):
            slide_parts = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_parts.append(shape.text.strip())
            if slide_parts:
                parts.append(f"[Slide {slide_num}]\n" + "\n".join(slide_parts))
        return "\n\n".join(parts)
    except Exception as exc:
        logger.warning("PPTX extraction failed: %s", exc)
        return ""


@router.post("/extract-text", response_model=ExtractResponse)
async def extract_text(
    file: UploadFile,
    current_user: User = Depends(get_current_user),
):
    """Extract plain text from a PDF, DOCX, or PPTX upload.

    Returns up to settings.BRIEF_MAX_CHARS characters (~450,000) of extracted
    text. The `truncated` flag indicates when the original text exceeded the cap.
    """
    filename = file.filename or "unknown"
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

    if ext not in ("pdf", "docx", "pptx"):
        raise HTTPException(
            status_code=status.HTTP_415_UNSUPPORTED_MEDIA_TYPE,
            detail=f"Unsupported file type: .{ext}. Supported: .pdf, .docx, .pptx",
        )

    data = await file.read()
    if len(data) > _MAX_FILE_BYTES:
        raise HTTPException(
            status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
            detail=f"File too large ({len(data) // 1024} KB). Maximum: {_MAX_FILE_BYTES // 1024} KB",
        )

    if ext == "pdf":
        text = _extract_pdf(data)
    elif ext == "docx":
        text = _extract_docx(data)
    else:
        text = _extract_pptx(data)

    if not text.strip():
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=f"Could not extract text from {filename}. The file may be image-based or encrypted.",
        )

    logger.info(
        "File extracted: user=%s filename=%s ext=%s chars=%d",
        current_user.id, filename, ext, len(text),
    )

    truncated = len(text) > _MAX_TEXT_CHARS
    return ExtractResponse(
        filename=filename,
        text=text[:_MAX_TEXT_CHARS],
        truncated=truncated,
    )
