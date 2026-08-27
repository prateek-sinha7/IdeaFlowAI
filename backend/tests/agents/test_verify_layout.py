"""tests/agents/test_verify_layout.py — the pptx shippability gate.

The deck that prompted this shipped as "1 violation" from a validator that
checked only shape geometry, while every headline on it printed through its own
body copy. These cases pin the three things that were wrong:

  * a 10x5.625 deck is pptxgenjs's DEFAULT 16:9 layout, not a defect — the old
    hard-coded 13.333x7.5 assertion cried wolf on the first line it printed, and
    the agent reading it dismissed the whole tool;
  * text that does not fit its box IS a defect, and it is the one that was
    invisible — python-pptx reports the box, never the rendered text;
  * an exact line spacing shorter than the font's line box guarantees overlapping
    lines, which is what "the pptx has overlapping text" nearly always is.

Decks are built here with python-pptx rather than fixtures, so each case states
its own geometry and nothing is inherited from a file nobody re-reads.
"""

from __future__ import annotations

import sys
from pathlib import Path

import pytest

pytest.importorskip("pptx")

from pptx import Presentation  # noqa: E402
from pptx.enum.text import MSO_AUTO_SIZE  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402

_SCRIPTS = (
    Path(__file__).resolve().parents[3]
    / "skills" / "opendesign" / "skills" / "pptx-html-fidelity-audit" / "scripts"
)
sys.path.insert(0, str(_SCRIPTS))

import verify_layout as V  # noqa: E402


def _deck(canvas=(13.333, 7.5)):
    prs = Presentation()
    prs.slide_width = Inches(canvas[0])
    prs.slide_height = Inches(canvas[1])
    prs.slides.add_slide(prs.slide_layouts[6])   # blank
    return prs


def _save(prs, tmp_path, name="deck.pptx") -> Path:
    path = tmp_path / name
    prs.save(str(path))
    return path


def _add_text(prs, text, *, left=0.5, top=0.5, width=6.0, height=1.0,
              size=18, font="Arial", bold=False, line_spacing=None):
    slide = prs.slides[0]
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    # python-pptx's add_textbox defaults to SHAPE_TO_FIT_TEXT. Say NONE
    # explicitly so each case states its own autofit rather than inheriting one.
    tf.auto_size = MSO_AUTO_SIZE.NONE
    p = tf.paragraphs[0]
    if line_spacing is not None:
        p.line_spacing = Pt(line_spacing)
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.name = font
    run.font.bold = bold
    return box


def _verify(path: Path):
    return V.verify(path, None, None, None, None)


# ════════════════════════════════════════════════════════════════════════════
# Canvas — the false alarm that got the whole tool ignored
# ════════════════════════════════════════════════════════════════════════════


class TestCanvas:
    def test_pptxgenjs_default_canvas_is_not_a_violation(self, tmp_path):
        """10x5.625 is LAYOUT_16x9, what you get when defineLayout is never
        called. Reporting it as a mismatch is what taught an agent that this tool
        is wrong about everything."""
        prs = _deck(canvas=(10.0, 5.625))
        _add_text(prs, "Fine", top=0.5, height=1.0, size=18)
        violations, _ = _verify(_save(prs, tmp_path))
        assert not [v for v in violations if "canvas" in v]

    def test_the_wide_canvas_is_not_a_violation_either(self, tmp_path):
        prs = _deck(canvas=(13.333, 7.5))
        _add_text(prs, "Fine")
        violations, _ = _verify(_save(prs, tmp_path))
        assert not [v for v in violations if "canvas" in v]

    def test_a_canvas_that_is_not_16_9_is_a_violation(self, tmp_path):
        prs = _deck(canvas=(10.0, 7.5))       # 4:3
        path = _save(prs, tmp_path)
        violations, _ = _verify(path)
        assert any("not 16:9" in v for v in violations)

    def test_an_explicitly_required_canvas_is_still_enforced(self, tmp_path):
        """A deck that MUST be a given size can still say so."""
        prs = _deck(canvas=(10.0, 5.625))
        path = _save(prs, tmp_path)
        violations, _ = V.verify(path, None, 13.333, 7.5, None)
        assert any("canvas mismatch" in v for v in violations)


# ════════════════════════════════════════════════════════════════════════════
# Text fit — the defect the geometry pass could not see
# ════════════════════════════════════════════════════════════════════════════


class TestTextFit:
    def test_text_taller_than_its_box_is_a_violation(self, tmp_path):
        prs = _deck()
        _add_text(
            prs,
            "Generic platforms were not designed for your regulatory environment",
            width=7.0, height=1.0, size=44,
        )
        violations, _ = _verify(_save(prs, tmp_path))
        assert any("text needs" in v and "tall" in v for v in violations)

    def test_text_that_fits_is_not_a_violation(self, tmp_path):
        prs = _deck()
        _add_text(prs, "Short heading", width=7.0, height=1.5, size=24)
        violations, _ = _verify(_save(prs, tmp_path))
        assert not [v for v in violations if "text needs" in v]

    def test_a_tiny_overflow_is_not_reported(self, tmp_path):
        """A page-number box a few hundredths of an inch over is a large RATIO of
        nothing at all. Reporting it is how a validator becomes noise and then
        gets ignored — which is the failure this whole file exists to prevent."""
        prs = _deck()
        # 9pt text needs ~0.15"; the inner height here is ~0.12" after insets, so
        # it is over by ~0.03" — under the absolute floor, above the ratio.
        _add_text(prs, "10 / 12", width=1.0, height=0.22, size=9)
        violations, _ = _verify(_save(prs, tmp_path))
        assert not [v for v in violations if "text needs" in v]

    def test_text_shrinking_autofit_is_honoured(self, tmp_path):
        """PowerPoint shrinks the text itself, so it is not overflowing."""
        prs = _deck()
        box = _add_text(prs, "A very long heading " * 8, width=6.0, height=0.6, size=40)
        box.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        violations, _ = _verify(_save(prs, tmp_path))
        assert not [v for v in violations if "text needs" in v]

    def test_an_unknown_font_is_still_measured(self, tmp_path):
        """A font nobody can resolve used to be skipped as UNVERIFIED — and a
        whole deck in that state exited 0 while every shape overflowed. It is now
        measured against a same-category stand-in (with extra slack, since the
        stand-in is not metric-exact), so the check runs on the deck that needs it
        most."""
        prs = _deck()
        _add_text(prs, "Overflowing " * 20, width=4.0, height=0.4, size=40,
                  font="No Such Font At All ZZZ")
        violations, notes = _verify(_save(prs, tmp_path))
        assert not [n for n in notes if "UNVERIFIED" in n]
        assert any("text needs" in v for v in violations)


# ════════════════════════════════════════════════════════════════════════════
# Fonts the recipient will not have
# ════════════════════════════════════════════════════════════════════════════


class TestFontSafety:
    def test_a_web_font_is_a_violation(self, tmp_path):
        """The source decks pull their type from Google Fonts, and transcribing
        those names into a .pptx produces a file that renders correctly on
        exactly one machine: the one that built it. PowerPoint stores a font
        NAME and substitutes whatever it likes when the name is missing —
        unless it's one render_pptx can embed (see test below); "Poppins" is
        deliberately NOT in the opendesign template font library."""
        prs = _deck()
        _add_text(prs, "Bespoke Employee Community", width=7.0, height=1.5,
                  size=24, font="Poppins")
        violations, _ = _verify(_save(prs, tmp_path))
        assert any("Poppins" in v and "PowerPoint ships with" in v for v in violations)

    def test_an_embeddable_web_font_is_not_a_violation(self, tmp_path):
        """DM Sans is a web font AND one render_pptx embeds real bytes for
        (skills/opendesign/fonts/) — the recipient never needs it installed,
        so this is not the same failure mode as test_a_web_font_is_a_violation."""
        prs = _deck()
        _add_text(prs, "Bespoke Employee Community", width=7.0, height=1.5,
                  size=24, font="DM Sans")
        violations, _ = _verify(_save(prs, tmp_path))
        assert not [v for v in violations if "PowerPoint ships with" in v]

    def test_office_fonts_are_not_flagged(self, tmp_path):
        prs = _deck()
        for i, fam in enumerate(["Arial", "Calibri", "Georgia", "Times New Roman",
                                 "Verdana", "Trebuchet MS", "Courier New"]):
            _add_text(prs, "Fine", left=0.5, top=0.4 + i * 0.6, width=3.0,
                      height=0.5, size=12, font=fam)
        violations, _ = _verify(_save(prs, tmp_path))
        assert not [v for v in violations if "PowerPoint ships with" in v]

    def test_one_line_per_font_not_per_shape(self, tmp_path):
        """125 identical lines is not 125 problems. A wall of them is how a
        validator gets skimmed instead of read."""
        prs = _deck()
        for i in range(6):
            _add_text(prs, f"Row {i}", left=0.5, top=0.4 + i * 0.7, width=3.0,
                      height=0.5, size=12, font="Poppins")
        violations, _ = _verify(_save(prs, tmp_path))
        font_lines = [v for v in violations if "PowerPoint ships with" in v]
        assert len(font_lines) == 1
        assert "6 shape(s)" in font_lines[0]

    def test_calibri_is_measurable_without_it_being_installed(self, tmp_path):
        """Calibri does not exist on macOS and Arial does not exist on a bare
        container. The bundled metric clones are what make the check work on any
        machine rather than degrading to UNVERIFIED for most of the deck."""
        prs = _deck()
        _add_text(prs, "Regulatory environment and compliance posture " * 4,
                  width=5.0, height=0.5, size=32, font="Calibri")
        violations, notes = _verify(_save(prs, tmp_path))
        assert not [n for n in notes if "UNVERIFIED" in n]
        assert any("text needs" in v for v in violations)


# ════════════════════════════════════════════════════════════════════════════
# Line spacing — what "overlapping text" actually is
# ════════════════════════════════════════════════════════════════════════════


class TestLineSpacing:
    def test_spacing_shorter_than_the_line_box_is_a_violation(self, tmp_path):
        """48pt text on 28pt exact leading — the real defect from run 62f84e1a.
        The box may be any size; the LINES collide with each other."""
        prs = _deck()
        _add_text(
            prs,
            "Generic platforms were not designed for your regulatory environment",
            width=7.0, height=4.0, size=48, font="Georgia", bold=True,
            line_spacing=28,
        )
        violations, _ = _verify(_save(prs, tmp_path))
        assert any("print through each other" in v for v in violations)

    def test_normal_spacing_is_not_a_violation(self, tmp_path):
        prs = _deck()
        _add_text(prs, "Two lines of perfectly ordinary body copy here",
                  width=3.0, height=3.0, size=18, line_spacing=24)
        violations, _ = _verify(_save(prs, tmp_path))
        assert not [v for v in violations if "print through" in v]

    def test_one_line_paragraphs_stacked_in_a_frame_still_collide(self, tmp_path):
        """The defect this guard was written wrong for.

        A title split as three ONE-LINE paragraphs — "Bespoke" / "Social" /
        "Platform", 72pt on 38pt leading — printed straight through itself, and
        the check reported nothing: it asked whether the PARAGRAPH had more than
        one line, when what collides is the FRAME's lines. Each paragraph had
        exactly one, so each was excused, and each printed over the next."""
        prs = _deck(canvas=(10.0, 5.625))
        slide = prs.slides[0]
        box = slide.shapes.add_textbox(Inches(0.5), Inches(1.1), Inches(4.5), Inches(2.4))
        tf = box.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        for i, word in enumerate(["Bespoke", "Social", "Platform"]):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.line_spacing = Pt(38)
            run = para.add_run()
            run.text = word
            run.font.size = Pt(72)
            run.font.name = "Arial"
            run.font.bold = True

        violations, _ = _verify(_save(prs, tmp_path))
        assert any("print through each other" in v for v in violations)
        assert any("3 lines" in v for v in violations)

    def test_a_single_line_cannot_overlap_itself(self, tmp_path):
        prs = _deck()
        _add_text(prs, "Short", width=7.0, height=2.0, size=40, line_spacing=12)
        violations, _ = _verify(_save(prs, tmp_path))
        assert not [v for v in violations if "print through" in v]


# ════════════════════════════════════════════════════════════════════════════
# Bounds and overlap
# ════════════════════════════════════════════════════════════════════════════


class TestBoundsAndOverlap:
    def test_a_shape_off_the_canvas_is_a_violation(self, tmp_path):
        prs = _deck(canvas=(10.0, 5.625))
        _add_text(prs, "Below the fold", top=5.4, height=0.6, size=12)
        violations, _ = _verify(_save(prs, tmp_path))
        assert any("exceeds canvas" in v for v in violations)

    def test_two_boxes_on_top_of_each_other_are_a_violation(self, tmp_path):
        prs = _deck()
        _add_text(prs, "Underneath", left=1.0, top=1.0, width=4.0, height=1.0, size=14)
        _add_text(prs, "On top", left=2.0, top=1.2, width=4.0, height=1.0, size=14)
        violations, _ = _verify(_save(prs, tmp_path))
        assert any("overlap by" in v for v in violations)

    def test_boxes_that_merely_touch_do_not_count(self, tmp_path):
        prs = _deck()
        _add_text(prs, "Above", left=1.0, top=1.0, width=4.0, height=1.0, size=14)
        _add_text(prs, "Below", left=1.0, top=2.0, width=4.0, height=1.0, size=14)
        violations, _ = _verify(_save(prs, tmp_path))
        assert not [v for v in violations if "overlap by" in v]


# ════════════════════════════════════════════════════════════════════════════
# The gate
# ════════════════════════════════════════════════════════════════════════════


class TestGate:
    def test_render_pptx_refuses_to_say_ok_for_a_broken_deck(self, tmp_path, monkeypatch):
        """The correction for how this failed: a verdict the agent has to fetch is
        a verdict it can dismiss, and one did. Verification lives inside the build
        so there is nothing to overrule."""
        from app.agents.tools import pptx_tools

        prs = _deck()
        _add_text(prs, "Generic platforms were not designed for your environment",
                  width=7.0, height=0.8, size=44)
        broken = _save(prs, tmp_path, "built.pptx").read_bytes()

        class _Sandbox:
            def path_for(self, name):
                return tmp_path / name

        monkeypatch.setattr(pptx_tools, "_SANDBOX", _Sandbox())
        monkeypatch.setattr(
            "app.services.pptx_export.generate_pptx_from_code",
            lambda code, title: broken,
        )

        out = pptx_tools.render_pptx.invoke({"pptxgenjs_code": "// whatever"})
        assert "NOT SHIPPABLE" in out
        assert "text needs" in out
        # The file is still written — a fixable deck on disk beats no deck.
        assert (tmp_path / pptx_tools.PPTX_NAME).is_file()

    def test_render_pptx_says_ok_for_a_clean_deck(self, tmp_path, monkeypatch):
        from app.agents.tools import pptx_tools

        prs = _deck()
        _add_text(prs, "Short heading", width=7.0, height=1.5, size=24)
        clean = _save(prs, tmp_path, "built.pptx").read_bytes()

        class _Sandbox:
            def path_for(self, name):
                return tmp_path / name

        monkeypatch.setattr(pptx_tools, "_SANDBOX", _Sandbox())
        monkeypatch.setattr(
            "app.services.pptx_export.generate_pptx_from_code",
            lambda code, title: clean,
        )

        out = pptx_tools.render_pptx.invoke({"pptxgenjs_code": "// whatever"})
        assert out.startswith("ok — wrote and verified")
